#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
故障诊断维修记录生成器 V2
生成5年的模拟故障维修记录数据并导出为Excel文件
包含数据故障代码映射
"""

import random
import pandas as pd
from datetime import datetime, timedelta
from dataclasses import dataclass
from typing import List, Dict, Tuple
import numpy as np

@dataclass
class MaintenanceMeasure:
    """维修措施数据类"""
    id: int
    operation: str
    duration: str
    tools: str

@dataclass
class FaultInfo:
    """故障信息数据类"""
    fault_type: str
    fault_code: str
    fault_name: str
    description: str
    severity: str
    priority: int
    maintenance_measures: List[MaintenanceMeasure]

# 数据故障类型定义
DATA_FAULT_TYPES = {
    '001': '满量程输出',
    '002': '零位输出',
    '003': '数据缺失',
    '004': '数据保持',
    '005': '剧烈波动异常',
    '006': '数据偏移',
    '006.01': '小幅正向偏移',
    '006.02': '小幅负向偏移',
    '006.03': '大幅正向偏移',
    '006.04': '大幅负向偏移',
    '006.05': '严重正向偏移',
    '006.06': '严重负向偏移'
}

# 仪表故障代码与数据故障类型的映射关系
FAULT_CODE_TO_DATA_FAULT = {
    # 数据偏差型故障主要对应数据偏移
    '302': ['006.01', '006.02'],  # 零点漂移50% - 小幅偏移
    '303': ['006.03', '006.04'],  # 零点漂移超限 - 大幅偏移
    '304': ['006.01', '006.02'],  # 灵敏度漂移50% - 小幅偏移
    '305': ['006.05', '006.06'],  # 灵敏度漂移超限 - 严重偏移
    '319': ['001', '002'],  # 磁力回路失衡 - 满量程/零位输出
    '320': ['006', '006.03'],  # 放大偏差过高 - 数据偏移
    '309-311': ['005'],  # 温度调节失效 - 剧烈波动
    
    # 数据传输中断型故障主要对应数据缺失、保持
    '101': ['003'],  # 系统控制器关停 - 数据缺失
    '116': ['003', '004'],  # Profibus错误 - 数据缺失/保持
    '201-209': ['003'],  # 总线中断 - 数据缺失
    '300': ['003', '004'],  # ADC无输出 - 数据缺失/保持
    '308': ['005', '003'],  # 计算错误 - 波动/缺失
    '318': ['003', '004'],  # ADC无新测量 - 数据缺失/保持
    '332-337': ['003', '004'],  # I/O板故障 - 数据缺失/保持
    '338-339': ['003', '001', '002'],  # 线路故障 - 缺失/满量程/零位
    
    # 数据保持型故障
    '301': ['001', '004'],  # 超出ADC阈值 - 满量程/数据保持
    '344': ['001'],  # 超上限130% - 满量程输出
    '345': ['002'],  # 低于下限-100% - 零位输出
    
    # 数据波动型故障
    '312': ['005', '006'],  # 压力修正失效 - 波动/偏移
    'EXT-01': ['005', '004'],  # 管路堵塞 - 波动/保持
    'EXT-02': ['005', '003']  # 气路泄漏 - 波动/缺失
}

def create_measure(id: int, operation: str, duration: str, tools: str) -> MaintenanceMeasure:
    """创建维修措施"""
    return MaintenanceMeasure(id, operation, duration, tools)

def initialize_fault_database() -> Dict[str, FaultInfo]:
    """初始化故障数据库"""
    db = {}
    
    # 数据偏差型故障
    db['302'] = FaultInfo(
        fault_type='数据偏差型故障',
        fault_code='302',
        fault_name='零点漂移超过允许范围50%',
        description='偏差漂移超过了允许范围的一半（±0.75vol%O2）',
        severity='中等',
        priority=2,
        maintenance_measures=[
            create_measure(1, '检查传感器表面污染', '10分钟', '清洁布、酒精'),
            create_measure(2, '检查采样管路冷凝水', '15分钟', '排水工具'),
            create_measure(3, '执行零点校准程序', '30分钟', '标准气体'),
            create_measure(4, '检查传感器老化程度', '1小时', '测试设备'),
            create_measure(5, '更换传感器模块', '2小时', '备用传感器')
        ]
    )
    
    db['303'] = FaultInfo(
        fault_type='数据偏差型故障',
        fault_code='303',
        fault_name='零点漂移超出允许范围',
        description='偏差漂移超出允许范围（±1.5vol%O2）',
        severity='严重',
        priority=1,
        maintenance_measures=[
            create_measure(1, '立即检查传感器状态灯', '5分钟', '无'),
            create_measure(2, '检查供电电压24V DC', '10分钟', '万用表'),
            create_measure(3, '紧急执行零点量程校准', '45分钟', '标准气体'),
            create_measure(4, '检查传感器参比电极', '1.5小时', '专用检测仪'),
            create_measure(5, '立即更换传感器', '2小时', '备用传感器')
        ]
    )
    
    db['304'] = FaultInfo(
        fault_type='数据偏差型故障',
        fault_code='304',
        fault_name='灵敏度漂移超出50%',
        description='放大漂移超出允许范围的50%',
        severity='中等',
        priority=2,
        maintenance_measures=[
            create_measure(1, '检查测量池窗片清洁度', '15分钟', '清洁工具'),
            create_measure(2, '检查光源强度稳定性', '20分钟', '光强测试仪'),
            create_measure(3, '执行量程标定', '40分钟', '量程气体'),
            create_measure(4, '检查检测器响应曲线', '1小时', '测试设备'),
            create_measure(5, '更换检测器组件', '2小时', '备用检测器')
        ]
    )
    
    db['305'] = FaultInfo(
        fault_type='数据偏差型故障',
        fault_code='305',
        fault_name='灵敏度漂移超出允许范围',
        description='放大漂移超出允许范围',
        severity='严重',
        priority=1,
        maintenance_measures=[
            create_measure(1, '停止测量检查报警', '5分钟', '无'),
            create_measure(2, '清洁所有光学元件', '30分钟', '光学清洁套装'),
            create_measure(3, '执行完整系统标定', '1小时', '多种标准气体'),
            create_measure(4, '调整光路对准', '1.5小时', '光路调整工具'),
            create_measure(5, '更换光源和检测器', '3小时', '备件')
        ]
    )
    
    db['319'] = FaultInfo(
        fault_type='数据偏差型故障',
        fault_code='319',
        fault_name='磁力测量回路失衡',
        description='磁力式传感器测量回路信号失去平衡',
        severity='严重',
        priority=1,
        maintenance_measures=[
            create_measure(1, '检查磁场线圈连接', '10分钟', '万用表'),
            create_measure(2, '测量磁力传感器输出', '20分钟', '示波器'),
            create_measure(3, '调节信号调理电路', '30分钟', '调试设备'),
            create_measure(4, '更换磁力传感器', '1小时', '备用传感器'),
            create_measure(5, '更换传感器电路板', '2小时', '备用电路板')
        ]
    )
    
    db['320'] = FaultInfo(
        fault_type='数据偏差型故障',
        fault_code='320',
        fault_name='测定放大偏差过高',
        description='信号放大器偏差超出正常范围',
        severity='中等',
        priority=2,
        maintenance_measures=[
            create_measure(1, '检查放大器供电', '10分钟', '万用表'),
            create_measure(2, '调整放大器零点增益', '25分钟', '示波器'),
            create_measure(3, '更换运放芯片', '45分钟', '备用芯片'),
            create_measure(4, '检查信号链路', '1小时', '信号发生器'),
            create_measure(5, '更换放大器板', '1.5小时', '备用电路板')
        ]
    )
    
    db['309-311'] = FaultInfo(
        fault_type='数据偏差型故障',
        fault_code='309-311',
        fault_name='温度调节器失效',
        description='温度控制超出范围',
        severity='中等',
        priority=2,
        maintenance_measures=[
            create_measure(1, '检查温度传感器', '10分钟', '温度计'),
            create_measure(2, '检查加热冷却器', '20分钟', '万用表'),
            create_measure(3, '校准PID参数', '40分钟', '调试软件'),
            create_measure(4, '更换温度传感器', '1小时', '备用传感器'),
            create_measure(5, '更换温控模块', '2小时', '备用模块')
        ]
    )
    
    # 数据传输中断型故障
    db['101'] = FaultInfo(
        fault_type='数据传输中断型故障',
        fault_code='101',
        fault_name='系统控制器关停',
        description='主控制器停止工作',
        severity='紧急',
        priority=1,
        maintenance_measures=[
            create_measure(1, '检查主电源保险丝', '5分钟', '万用表'),
            create_measure(2, '检查24V电源输出', '10分钟', '万用表'),
            create_measure(3, '重启控制器系统', '15分钟', '无'),
            create_measure(4, '检查CPU和内存', '30分钟', '诊断软件'),
            create_measure(5, '更换控制器主板', '2小时', '备用主板')
        ]
    )
    
    db['116'] = FaultInfo(
        fault_type='数据传输中断型故障',
        fault_code='116',
        fault_name='Profibus安装错误',
        description='Profibus模块安装位置错误',
        severity='中等',
        priority=3,
        maintenance_measures=[
            create_measure(1, '确认安装位置', '5分钟', '无'),
            create_measure(2, '重装到X20/X21槽', '15分钟', '螺丝刀'),
            create_measure(3, '重配通讯参数', '20分钟', '配置软件'),
            create_measure(4, '测试通讯连接', '30分钟', 'Profibus测试仪'),
            create_measure(5, '更换Profibus模块', '1小时', '备用模块')
        ]
    )
    
    db['201-209'] = FaultInfo(
        fault_type='数据传输中断型故障',
        fault_code='201-209',
        fault_name='系统总线连接中断',
        description='系统总线通讯中断',
        severity='严重',
        priority=1,
        maintenance_measures=[
            create_measure(1, '检查总线电缆', '10分钟', '无'),
            create_measure(2, '检查终端电阻', '15分钟', '万用表'),
            create_measure(3, '更换总线电缆', '30分钟', '备用电缆'),
            create_measure(4, '检查模块供电', '20分钟', '万用表'),
            create_measure(5, '更换通讯模块', '1小时', '备用模块')
        ]
    )
    
    db['300'] = FaultInfo(
        fault_type='数据传输中断型故障',
        fault_code='300',
        fault_name='模数转换器无输出',
        description='ADC无新测量值',
        severity='严重',
        priority=1,
        maintenance_measures=[
            create_measure(1, '检查ADC供电', '10分钟', '万用表'),
            create_measure(2, '检查模拟输入', '15分钟', '示波器'),
            create_measure(3, '重置ADC芯片', '20分钟', '复位工具'),
            create_measure(4, '更换ADC芯片', '1小时', '备用芯片'),
            create_measure(5, '更换采集卡', '1.5小时', '备用采集卡')
        ]
    )
    
    db['308'] = FaultInfo(
        fault_type='数据传输中断型故障',
        fault_code='308',
        fault_name='测定值计算错误',
        description='计算过程错误',
        severity='中等',
        priority=2,
        maintenance_measures=[
            create_measure(1, '重启处理程序', '5分钟', '无'),
            create_measure(2, '检查CPU内存', '10分钟', '监控软件'),
            create_measure(3, '清理系统缓存', '15分钟', '清理工具'),
            create_measure(4, '重装固件程序', '45分钟', '固件包'),
            create_measure(5, '更换处理器板', '2小时', '备用板')
        ]
    )
    
    db['318'] = FaultInfo(
        fault_type='数据传输中断型故障',
        fault_code='318',
        fault_name='ADC无新测量',
        description='ADC无新数据',
        severity='严重',
        priority=1,
        maintenance_measures=[
            create_measure(1, '检查触发信号', '10分钟', '示波器'),
            create_measure(2, '检查时钟信号', '15分钟', '示波器'),
            create_measure(3, '重配采样参数', '20分钟', '配置软件'),
            create_measure(4, '更换时钟芯片', '45分钟', '备用芯片'),
            create_measure(5, '更换ADC模块', '1.5小时', '备用模块')
        ]
    )
    
    db['332-337'] = FaultInfo(
        fault_type='数据传输中断型故障',
        fault_code='332-337',
        fault_name='I/O板故障',
        description='I/O板硬件问题',
        severity='中等',
        priority=2,
        maintenance_measures=[
            create_measure(1, '检查I/O板指示灯', '5分钟', '无'),
            create_measure(2, '检查I/O配置', '15分钟', '配置软件'),
            create_measure(3, '测试I/O通道', '30分钟', '万用表'),
            create_measure(4, '初始化I/O板', '20分钟', '初始化工具'),
            create_measure(5, '更换I/O板', '1小时', '备用板')
        ]
    )
    
    db['338-339'] = FaultInfo(
        fault_type='数据传输中断型故障',
        fault_code='338-339',
        fault_name='模拟线路故障',
        description='线路断裂或短路',
        severity='严重',
        priority=1,
        maintenance_measures=[
            create_measure(1, '检查接线端子', '10分钟', '螺丝刀'),
            create_measure(2, '测量线路通断', '15分钟', '万用表'),
            create_measure(3, '检查屏蔽接地', '20分钟', '接地测试仪'),
            create_measure(4, '更换信号电缆', '30分钟', '备用电缆'),
            create_measure(5, '重新布线', '2小时', '布线工具')
        ]
    )
    
    # 数据保持型故障
    db['301'] = FaultInfo(
        fault_type='数据保持型故障',
        fault_code='301',
        fault_name='超出ADC阈值',
        description='超出ADC范围',
        severity='中等',
        priority=2,
        maintenance_measures=[
            create_measure(1, '确认实际浓度', '5分钟', '便携式分析仪'),
            create_measure(2, '检查量程设置', '10分钟', '配置软件'),
            create_measure(3, '调整信号衰减', '20分钟', '调节工具'),
            create_measure(4, '重选测量量程', '30分钟', '配置软件'),
            create_measure(5, '更换大量程传感器', '2小时', '备用传感器')
        ]
    )
    
    db['344'] = FaultInfo(
        fault_type='数据保持型故障',
        fault_code='344',
        fault_name='超上限130%',
        description='超过量程130%',
        severity='严重',
        priority=1,
        maintenance_measures=[
            create_measure(1, '检查工艺异常', '5分钟', '工艺参数表'),
            create_measure(2, '检查空气泄漏', '15分钟', '检漏仪'),
            create_measure(3, '检查传感器饱和', '20分钟', '测试仪'),
            create_measure(4, '切换高量程', '30分钟', '配置软件'),
            create_measure(5, '更换高量程传感器', '2小时', '备用传感器')
        ]
    )
    
    db['345'] = FaultInfo(
        fault_type='数据保持型故障',
        fault_code='345',
        fault_name='低于下限-100%',
        description='低于量程-100%',
        severity='严重',
        priority=1,
        maintenance_measures=[
            create_measure(1, '检查传感器接线', '10分钟', '接线图'),
            create_measure(2, '检查信号电路', '20分钟', '万用表'),
            create_measure(3, '验证零点设置', '30分钟', '零点气体'),
            create_measure(4, '重新标定', '45分钟', '标准气体'),
            create_measure(5, '更换传感器信号板', '2小时', '备件')
        ]
    )
    
    # 数据波动型故障
    db['312'] = FaultInfo(
        fault_type='数据波动型故障',
        fault_code='312',
        fault_name='压力修正失效',
        description='压力测量错误',
        severity='中等',
        priority=2,
        maintenance_measures=[
            create_measure(1, '检查压力传感器', '10分钟', '压力表'),
            create_measure(2, '检查压力管路', '15分钟', '检漏仪'),
            create_measure(3, '校准压力传感器', '30分钟', '标准压力源'),
            create_measure(4, '检查补偿算法', '20分钟', '配置软件'),
            create_measure(5, '更换压力传感器', '1小时', '备用传感器')
        ]
    )
    
    db['EXT-01'] = FaultInfo(
        fault_type='数据波动型故障',
        fault_code='EXT-01',
        fault_name='管路污染堵塞',
        description='管道或过滤器问题',
        severity='中等',
        priority=2,
        maintenance_measures=[
            create_measure(1, '检查过滤器压差', '5分钟', '压差表'),
            create_measure(2, '更换过滤器', '20分钟', '备用滤芯'),
            create_measure(3, '吹扫采样管路', '30分钟', '压缩空气'),
            create_measure(4, '检查管路泄漏', '45分钟', '检漏仪'),
            create_measure(5, '更换采样管路', '2小时', '备用管路')
        ]
    )
    
    db['EXT-02'] = FaultInfo(
        fault_type='数据波动型故障',
        fault_code='EXT-02',
        fault_name='气路扭结泄漏',
        description='气路系统问题',
        severity='中等',
        priority=2,
        maintenance_measures=[
            create_measure(1, '检查管路扭结', '5分钟', '手电筒'),
            create_measure(2, '检查流量', '10分钟', '流量计'),
            create_measure(3, '逐段检漏', '30分钟', '检漏仪'),
            create_measure(4, '紧固接头', '20分钟', '扳手'),
            create_measure(5, '重装气路', '3小时', '全套管路')
        ]
    )
    
    return db

def get_data_fault_code(instrument_fault_code: str) -> str:
    """根据仪表故障代码获取对应的数据故障代码"""
    if instrument_fault_code in FAULT_CODE_TO_DATA_FAULT:
        data_fault_codes = FAULT_CODE_TO_DATA_FAULT[instrument_fault_code]
        selected_code = random.choice(data_fault_codes)
        return f"{selected_code} {DATA_FAULT_TYPES[selected_code]}"
    else:
        # 默认返回数据偏移
        return "006 数据偏移"

def generate_fault_records(start_date: datetime, end_date: datetime, num_records: int) -> List[Dict]:
    """生成故障记录"""
    # 初始化故障数据库
    fault_db = initialize_fault_database()
    fault_codes = list(fault_db.keys())
    
    # 生成记录
    records = []
    
    # 计算时间范围
    time_range = (end_date - start_date).total_seconds()
    
    for i in range(num_records):
        # 随机选择故障代码
        fault_code = random.choice(fault_codes)
        fault_info = fault_db[fault_code]
        
        # 随机选择维修措施
        measure = random.choice(fault_info.maintenance_measures)
        
        # 生成随机时间
        random_seconds = random.uniform(0, time_range)
        fault_time = start_date + timedelta(seconds=random_seconds)
        
        # 获取对应的数据故障代码
        data_fault_code = get_data_fault_code(fault_code)
        
        # 创建记录
        record = {
            '数据故障代码': data_fault_code,
            '故障时间': fault_time.strftime('%Y-%m-%d %H:%M:%S'),
            '仪表故障代码': fault_code,
            '故障描述/名称': fault_info.fault_name,
            '故障类型': fault_info.fault_type,
            '严重程度': fault_info.severity,
            '优先级': fault_info.priority,
            '维修操作': measure.operation,
            '耗时': measure.duration,
            '工具': measure.tools
        }
        
        records.append(record)
    
    # 按时间排序
    records.sort(key=lambda x: x['故障时间'])
    
    return records

def export_to_excel(records: List[Dict], filename: str):
    """导出记录到Excel文件"""
    # 创建DataFrame
    df = pd.DataFrame(records)
    
    # 创建Excel写入器
    with pd.ExcelWriter(filename, engine='openpyxl') as writer:
        # 写入数据
        df.to_excel(writer, sheet_name='故障维修记录', index=False)
        
        # 获取工作表
        worksheet = writer.sheets['故障维修记录']
        
        # 设置列宽
        column_widths = {
            'A': 25,  # 数据故障代码
            'B': 20,  # 故障时间
            'C': 15,  # 仪表故障代码
            'D': 30,  # 故障描述/名称
            'E': 20,  # 故障类型
            'F': 10,  # 严重程度
            'G': 10,  # 优先级
            'H': 30,  # 维修操作
            'I': 15,  # 耗时
            'J': 25   # 工具
        }
        
        for col, width in column_widths.items():
            worksheet.column_dimensions[col].width = width
        
        # 添加标题样式
        from openpyxl.styles import Font, PatternFill, Alignment
        
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center')
        
        for cell in worksheet[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
        
        # 添加数据行样式
        data_alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        
        for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
            for cell in row:
                cell.alignment = data_alignment
                
                # 根据严重程度设置颜色
                if cell.column == 6:  # 严重程度列（现在是第6列）
                    if cell.value == '紧急':
                        cell.fill = PatternFill(start_color='FF0000', end_color='FF0000', fill_type='solid')
                        cell.font = Font(color='FFFFFF', bold=True)
                    elif cell.value == '严重':
                        cell.fill = PatternFill(start_color='FF9900', end_color='FF9900', fill_type='solid')
                        cell.font = Font(bold=True)
                    elif cell.value == '中等':
                        cell.fill = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')
        
        # 冻结首行
        worksheet.freeze_panes = 'A2'
        
        # 添加筛选
        worksheet.auto_filter.ref = worksheet.dimensions
    
    print(f"Excel文件已生成: {filename}")
    return filename

def generate_distribution_stats(records: List[Dict]):
    """生成分布统计信息"""
    df = pd.DataFrame(records)
    
    print("\n=== 故障记录统计信息 ===")
    print(f"总记录数: {len(records)}")
    print(f"\n故障类型分布:")
    print(df['故障类型'].value_counts())
    print(f"\n严重程度分布:")
    print(df['严重程度'].value_counts())
    print(f"\n故障代码分布 (前10):")
    print(df['仪表故障代码'].value_counts().head(10))
    print(f"\n数据故障类型分布:")
    # 提取数据故障代码的主要类型
    df['数据故障类型主码'] = df['数据故障代码'].str.split(' ').str[0].str.split('.').str[0]
    print(df['数据故障类型主码'].value_counts())
    
    # 计算平均每月故障数
    df['故障时间'] = pd.to_datetime(df['故障时间'])
    monthly_counts = df.groupby(df['故障时间'].dt.to_period('M')).size()
    print(f"\n平均每月故障数: {monthly_counts.mean():.1f}")

def main():
    """主函数"""
    # 设置时间范围（5年）
    end_date = datetime.now()
    start_date = end_date - timedelta(days=5*365)
    
    # 生成记录数量：固定60条
    num_records = 60
    
    print(f"开始生成故障记录...")
    print(f"时间范围: {start_date.strftime('%Y-%m-%d')} 至 {end_date.strftime('%Y-%m-%d')}")
    print(f"预计生成记录数: {num_records}")
    
    # 生成故障记录
    records = generate_fault_records(start_date, end_date, num_records)
    
    # 导出到Excel
    filename = f"故障诊断维修记录_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    export_to_excel(records, filename)
    
    # 显示统计信息
    generate_distribution_stats(records)
    
    print(f"\n任务完成！文件保存为: {filename}")

if __name__ == "__main__":
    main()